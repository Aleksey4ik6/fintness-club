import csv
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


DISPLAY_VALUES = {
    "active": "Активен",
    "inactive": "Неактивен",
    "expired": "Истек",
    "closed": "Закрыт",
    "cash": "Наличные",
    "card": "Карта",
    "transfer": "Перевод",
}


class ExportService:
    def export_csv(self, path: str, headers: list[str], rows: list[dict[str, Any]], keys: list[str]) -> None:
        with Path(path).open("w", encoding="utf-8-sig", newline="") as file:
            writer = csv.writer(file, delimiter=";")
            writer.writerow(headers)
            for row in rows:
                writer.writerow([self._format_value(row.get(key)) for key in keys])

    def export_xlsx(self, path: str, title: str, headers: list[str], rows: list[dict[str, Any]], keys: list[str]) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = title[:31]

        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F7A5C")

        for row in rows:
            sheet.append([self._format_value(row.get(key)) for key in keys])

        for column in sheet.columns:
            width = max(len(str(cell.value or "")) for cell in column) + 2
            sheet.column_dimensions[column[0].column_letter].width = min(width, 42)

        workbook.save(path)

    def _format_value(self, value: Any) -> str:
        if value is None:
            return ""
        return DISPLAY_VALUES.get(str(value), str(value))
